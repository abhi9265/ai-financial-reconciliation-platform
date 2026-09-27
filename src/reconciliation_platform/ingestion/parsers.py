"""Source-specific parsers that convert external files into contract-shaped rows.

Parsers stop at a stable tabular representation. Normalization into the canonical
transaction model remains a separate concern so source parsing cannot silently
change reconciliation semantics.
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from typing import Any, Mapping

from reconciliation_platform.ingestion.contracts import (
    IngestionContractError,
    normalize_column_name,
    validate_columns,
    validate_rows,
)
from reconciliation_platform.models.canonical_transaction import SourceSystem


class SourceParseError(ValueError):
    """Raised when an uploaded source file cannot be parsed safely."""


@dataclass(frozen=True)
class ParsedSource:
    source_system: SourceSystem
    filename: str
    columns: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]


# External exports frequently rename the same business field. These aliases are
# intentionally source-specific rather than a global fuzzy column mapper.
_COLUMN_ALIASES: Mapping[SourceSystem, Mapping[str, str]] = {
    SourceSystem.BANK: {
        "txn_id": "transaction_id",
        "transaction_id": "transaction_id",
        "transaction_no": "transaction_id",
        "transaction_number": "transaction_id",
        "date": "transaction_date",
        "txn_date": "transaction_date",
        "transaction_date": "transaction_date",
        "value_date": "value_date",
        "debit": "debit",
        "withdrawal": "debit",
        "debit_amount": "debit",
        "credit": "credit",
        "deposit": "credit",
        "credit_amount": "credit",
        "amount": "amount",
        "transaction_amount": "amount",
        "narration": "narration",
        "description": "narration",
        "remarks": "narration",
        "reference": "reference",
        "reference_no": "reference",
        "ref_no": "reference",
        "account_number": "account_number",
        "account_no": "account_number",
    },
    SourceSystem.PURCHASE_REGISTER: {
        "invoice_record_id": "invoice_record_id",
        "record_id": "invoice_record_id",
        "invoice_id": "invoice_record_id",
        "invoice_number": "invoice_number",
        "invoice_no": "invoice_number",
        "bill_no": "invoice_number",
        "invoice_date": "invoice_date",
        "bill_date": "invoice_date",
        "vendor_name": "vendor_name",
        "supplier_name": "vendor_name",
        "vendor": "vendor_name",
        "gstin": "gstin",
        "vendor_gstin": "gstin",
        "supplier_gstin": "gstin",
        "taxable_value": "taxable_value",
        "taxable_amount": "taxable_value",
        "taxable": "taxable_value",
        "cgst": "cgst",
        "sgst": "sgst",
        "igst": "igst",
        "total": "total",
        "total_amount": "total",
        "invoice_total": "total",
    },
    SourceSystem.TALLY: {
        "transaction_date": "transaction_date",
        "date": "transaction_date",
        "amount": "amount",
        "transaction_type": "transaction_type",
        "type": "transaction_type",
        "source_record_id": "source_record_id",
        "record_id": "source_record_id",
        "transaction_id": "source_record_id",
        "reference_number": "reference_number",
        "reference_no": "reference_number",
        "account_name": "account_name",
        "ledger": "account_name",
        "description": "description",
        "narration": "description",
        "debit": "debit",
        "credit": "credit",
    },
}


def _canonical_header(value: str) -> str:
    try:
        return normalize_column_name(value)
    except IngestionContractError as exc:
        raise SourceParseError(str(exc)) from exc


def _map_headers(source_system: SourceSystem, headers: list[str]) -> tuple[str, ...]:
    aliases = _COLUMN_ALIASES.get(source_system, {})
    mapped: list[str] = []
    for header in headers:
        normalized = _canonical_header(header)
        mapped.append(aliases.get(normalized, normalized))

    if len(mapped) != len(set(mapped)):
        duplicates = sorted({name for name in mapped if mapped.count(name) > 1})
        raise SourceParseError(
            f"source headers map to duplicate canonical columns: {', '.join(duplicates)}"
        )
    return tuple(mapped)


def _parse_csv(content: bytes, source_system: SourceSystem, filename: str) -> ParsedSource:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise SourceParseError("CSV must be UTF-8 encoded") from exc

    try:
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise SourceParseError("CSV must contain a header row")
        columns = _map_headers(source_system, list(reader.fieldnames))
        rows: list[dict[str, Any]] = []
        for row_number, raw in enumerate(reader, start=2):
            if None in raw:
                raise SourceParseError(f"row {row_number} contains more values than headers")
            rows.append(
                {
                    columns[index]: value.strip() if isinstance(value, str) else value
                    for index, (_, value) in enumerate(raw.items())
                }
            )
    except csv.Error as exc:
        raise SourceParseError(f"invalid CSV: {exc}") from exc

    if not rows:
        raise SourceParseError("source file contains no data rows")

    # Many bank exports expose debit/credit but omit a separate amount column.
    # Derive the canonical amount before applying the strict source contract.
    if source_system is SourceSystem.BANK and "amount" not in columns:
        if "debit" not in columns and "credit" not in columns:
            raise SourceParseError("bank source requires amount or debit/credit columns")
        rows = [
            {
                **row,
                "amount": (
                    row.get("debit")
                    if str(row.get("debit") or "").strip()
                    else row.get("credit")
                ),
            }
            for row in rows
        ]
        columns = tuple((*columns, "amount"))

    try:
        validate_columns(source_system, columns)
        validate_rows(rows, columns)
    except IngestionContractError as exc:
        raise SourceParseError(str(exc)) from exc

    return ParsedSource(source_system, filename, columns, tuple(rows))


def _records_from_json(payload: Any) -> list[Mapping[str, Any]]:
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, Mapping)]
    if isinstance(payload, Mapping):
        for key in ("invoices", "records", "data", "items"):
            candidate = payload.get(key)
            if isinstance(candidate, list):
                return [item for item in candidate if isinstance(item, Mapping)]
        return [payload]
    raise SourceParseError("JSON source must contain an object or array of records")


def _flatten_json_record(record: Mapping[str, Any]) -> dict[str, Any]:
    flattened: dict[str, Any] = {}
    for key, value in record.items():
        if isinstance(value, Mapping):
            for nested_key, nested_value in value.items():
                flattened[f"{key}_{nested_key}"] = nested_value
        else:
            flattened[key] = value
    return flattened


def _parse_xlsx(content: bytes, source_system: SourceSystem, filename: str) -> ParsedSource:
    if source_system is not SourceSystem.TALLY:
        raise SourceParseError("XLSX uploads currently support the Tally source only")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise SourceParseError("invalid XLSX workbook") from exc

    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        header_row: tuple[Any, ...] | None = None
        buffered_rows: list[tuple[Any, ...]] = []
        for _ in range(20):
            try:
                row = next(rows_iter)
            except StopIteration:
                break
            if sum(value is not None and str(value).strip() != "" for value in row) >= 2:
                header_row = row
                break
        if header_row is None:
            raise SourceParseError("XLSX does not contain a usable header row")

        raw_headers = [str(value) for value in header_row]
        columns = _map_headers(source_system, raw_headers)
        for row in rows_iter:
            values = list(row[: len(columns)])
            if not any(value is not None and str(value).strip() != "" for value in values):
                continue
            values.extend([None] * (len(columns) - len(values)))
            buffered_rows.append(
                {
                    columns[index]: value.strip() if isinstance(value, str) else value
                    for index, value in enumerate(values)
                }
            )
    except SourceParseError:
        raise
    except Exception as exc:
        raise SourceParseError("failed to read XLSX rows") from exc
    finally:
        workbook.close()

    if not buffered_rows:
        raise SourceParseError("XLSX contains no data rows")
    try:
        validate_columns(source_system, columns)
        validate_rows(buffered_rows, columns)
    except IngestionContractError as exc:
        raise SourceParseError(str(exc)) from exc
    return ParsedSource(source_system, filename, columns, tuple(buffered_rows))


def _parse_gst_json(content: bytes, filename: str) -> ParsedSource:
    try:
        payload = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise SourceParseError("invalid UTF-8 JSON") from exc

    records = _records_from_json(payload)
    if not records:
        raise SourceParseError("JSON source contains no data records")

    aliases = {
        "invoice_number": "invoice_number",
        "invoice_no": "invoice_number",
        "invoice_date": "invoice_date",
        "date": "invoice_date",
        "amount": "amount",
        "total": "amount",
        "invoice_total": "amount",
        "counterparty_gstin": "counterparty_gstin",
        "gstin": "counterparty_gstin",
        "buyer_gstin": "counterparty_gstin",
        "supplier_gstin": "counterparty_gstin",
        "source_record_id": "source_record_id",
        "record_id": "source_record_id",
        "invoice_id": "source_record_id",
        "counterparty_name": "counterparty_name",
        "buyer_name": "counterparty_name",
        "supplier_name": "counterparty_name",
        "taxable_amount": "taxable_amount",
        "taxable_value": "taxable_amount",
        "cgst": "cgst",
        "sgst": "sgst",
        "igst": "igst",
        "total_tax": "total_tax",
    }

    rows: list[dict[str, Any]] = []
    for record in records:
        flattened = _flatten_json_record(record)
        mapped: dict[str, Any] = {}
        for key, value in flattened.items():
            normalized = _canonical_header(key)
            mapped[aliases.get(normalized, normalized)] = value
        if "source_record_id" not in mapped and "invoice_number" in mapped:
            mapped["source_record_id"] = str(mapped["invoice_number"])
        rows.append(mapped)

    columns = tuple(dict.fromkeys(key for row in rows for key in row))
    try:
        validate_columns(SourceSystem.GST, columns)
        validate_rows(rows, columns)
    except IngestionContractError as exc:
        raise SourceParseError(str(exc)) from exc

    return ParsedSource(SourceSystem.GST, filename, columns, tuple(rows))


def parse_source(content: bytes, *, source_system: SourceSystem, filename: str) -> ParsedSource:
    """Parse a supported uploaded source into validated contract-shaped rows."""
    normalized_name = re.sub(r"[^a-zA-Z0-9._-]", "_", filename or "upload")
    suffix = normalized_name.rsplit(".", 1)[-1].lower() if "." in normalized_name else ""
    if suffix == "csv":
        return _parse_csv(content, source_system, normalized_name)
    if suffix == "xlsx":
        return _parse_xlsx(content, source_system, normalized_name)
    if suffix == "json" and source_system is SourceSystem.GST:
        return _parse_gst_json(content, normalized_name)
    raise SourceParseError(
        f"unsupported file format '{suffix or 'unknown'}' for source '{source_system.value}'"
    )
